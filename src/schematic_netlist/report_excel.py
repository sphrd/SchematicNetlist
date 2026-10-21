"""Analysis workbook: Summary, Components, Pins, Nets and Warnings sheets.

Written in openpyxl's write-only mode: a large schematic gives tens of thousands of pin rows.
"""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .circuit import Circuit

_FILL = {
    "rail": "D9D9D9",
    "warn": "F8CBAD",
    "info": "FFF2CC",
    "passive": "E2EFDA",
    "IC": "DDEBF7",
    "relay": "FFE699",
    "connector": "EDEDED",
    "unknown": "F4CCCC",
}
_KIND_FILL = {
    "resistor": "passive", "capacitor": "passive", "inductor": "passive", "ferrite bead": "passive",
    "resistor network": "passive", "thermistor/varistor": "passive", "crystal": "passive", "fuse": "passive",
    "jumper": "passive", "diode": "passive",
    "IC": "IC", "transistor (BJT)": "IC", "MOSFET": "IC",
    "relay": "relay", "relay contact": "relay", "relay coil": "relay", "switch": "relay",
    "connector": "connector", "test point": "connector",
    "unknown": "unknown",
}
_FILLS = {k: PatternFill(start_color=c, end_color=c, fill_type="solid") for k, c in _FILL.items()}
_BOLD = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")
_MAX_CONN = 2000  # chars of the Nets "Connections" column (Excel cells cap at 32767)


def _sheet(wb: Workbook, title: str, header: list[str], widths: list[int], rows: int):
    ws = wb.create_sheet(title)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{rows + 1}"
    cells = []
    for h in header:
        c = WriteOnlyCell(ws, value=h)
        c.font = _BOLD
        cells.append(c)
    ws.append(cells)
    return ws


def _row(ws, values: list, fills: dict[int, str] | None = None, wrap: bool = False) -> None:
    if not fills and not wrap:
        ws.append(values)
        return
    cells = []
    for i, v in enumerate(values):
        c = WriteOnlyCell(ws, value=v)
        if fills and i in fills:
            c.fill = _FILLS[fills[i]]
        if wrap:
            c.alignment = _WRAP
        cells.append(c)
    ws.append(cells)


def write_report(path: str | Path, circuit: Circuit, title: str) -> None:
    wb = Workbook(write_only=True)
    kinds = circuit.kinds()
    warn = sum(1 for w in circuit.warnings if w.level == "warn")
    info = len(circuit.warnings) - warn

    ws = _sheet(wb, "Summary", ["Item", "Value"], [28, 90], 0)
    rows = [
        ("Title", title),
        ("Source PDF", circuit.source),
        ("Pages read", len(circuit.pages)),
        ("Pages by source", ", ".join(f"{k}: {v}" for k, v in sorted(Counter(circuit.page_source.values()).items()))),
        ("Parts", len(circuit.parts)),
        ("Nets (with pins)", len(circuit.nets)),
        ("Power/ground rails", ", ".join(n.name for n in circuit.rails)),
        ("Warnings / notes", f"{warn} / {info}"),
        ("", ""),
        ("Parts by kind", ""),
    ]
    rows += [(f"  {k}", v) for k, v in kinds.most_common()]
    rows += [("", ""), ("Engine log", "")] + [("", line) for line in circuit.log]
    for r in rows:
        _row(ws, list(r), wrap=True)

    ws = _sheet(wb, "Components", ["Refdes", "Kind", "Value / Part", "Pins", "Pages", "Group", "Shape", "Notes"],
                [14, 20, 24, 6, 12, 12, 12, 60], len(circuit.parts))
    for p in circuit.parts:
        key = _KIND_FILL.get(p.kind)
        _row(ws, [p.refdes, p.kind, p.value, p.pin_count, ", ".join(map(str, p.pages)), p.group, p.shape,
                  "; ".join(p.notes)], {1: key} if key else None)

    rail_names = {n.name for n in circuit.rails}
    n_pins = sum(p.pin_count for p in circuit.parts)
    ws = _sheet(wb, "Pins", ["Refdes", "Kind", "Pin", "Pin name", "Net", "Rail", "Page"], [14, 20, 8, 18, 28, 6, 6], n_pins)
    for p in circuit.parts:
        for q in p.pins:
            rail = q.net in rail_names
            _row(ws, [p.refdes, p.kind, q.number, q.name, q.net, "yes" if rail else "", q.page], {4: "rail"} if rail else None)

    ws = _sheet(wb, "Nets", ["Net", "Rail", "Pins", "Pages", "Aliases", "Unnamed flags", "Connections"],
                [28, 8, 6, 14, 24, 8, 100], len(circuit.nets))
    for n in circuit.nets:
        conn = "  ".join(f"{r}.{p or '?'}" + (f"/{nm}" if nm else "") for r, p, nm in n.pins)
        if len(conn) > _MAX_CONN:
            conn = conn[:_MAX_CONN] + " ..."
        _row(ws, [n.name, n.rail_reason, len(n.pins), ", ".join(map(str, n.pages)), ", ".join(n.aliases[1:]),
                  n.flags or "", conn], {0: "rail"} if n.is_rail else None)

    ws = _sheet(wb, "Warnings", ["Level", "Page", "Message"], [8, 6, 120], len(circuit.warnings))
    for w in circuit.warnings:
        _row(ws, [w.level, w.page if w.page is not None else "", w.text], {0: w.level})

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
